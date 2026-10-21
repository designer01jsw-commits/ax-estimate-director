# -*- coding: utf-8 -*-
"""가상 샘플 엑셀 데이터 생성 스크립트"""

from pathlib import Path
import openpyxl

DATA_DIR = Path(__file__).resolve().parent.parent / "data"
SAMPLES_DIR = DATA_DIR / "samples"
BROKEN_DIR = SAMPLES_DIR / "broken"

def create_p001():
    # 1. P001 최초견적
    wb1 = openpyxl.Workbook()
    ws1 = wb1.active
    ws1.title = "견적"
    ws1.append(["SAMPLE — 데모용 가상 데이터"])
    ws1.append([])
    ws1.append(["항목코드", "구분", "항목명", "규격", "수량", "단위", "단가", "금액"])
    ws1.append(["E-001", "제작", "아크릴 케이스", "600x400x300 T5", 20, "EA", 85000, 1700000])
    ws1.append(["E-002", "제작", "목공 받침대", "자작나무 18T", 20, "EA", 120000, 2400000])
    ws1.append(["E-003", "인쇄", "그래픽 패널 출력", "UV 1200x900", 40, "EA", 45000, 1850000])  # 산출식 불일치 의도
    ws1.append(["E-004", "전기전자", "LED 조명 모듈", "12V 3000K", 40, "EA", 18000, 720000])
    ws1.append(["E-005", "전기전자", "터치 센서 키트", "정전식", 20, "SET", 65000, 1300000])
    ws1.append(["E-006", "운송", "운송비", "5톤 윙바디", 2, "회", 450000, 900000])
    ws1.append(["E-007", "디자인", "원형 디자인", "컨셉~CMF", 5, "일", 400000, 2000000])
    ws1.append(["E-010", "인쇄", "안내 리플릿 인쇄", "A4 3단", 1000, "부", 350, 350000])
    ws1.append(["", "", "합계", "", "", "", "", ""])

    ws_labor = wb1.create_sheet(title="인력")
    ws_labor.append(["SAMPLE — 데모용 가상 데이터"])
    ws_labor.append([])
    ws_labor.append(["항목코드", "공정", "역할", "인원", "일수", "일단가"])
    ws_labor.append(["L-001", "설치", "목공", 2, 3, 300000])
    ws_labor.append(["L-002", "설치", "전기", 1, 2, 320000])
    ws_labor.append(["L-003", "철거", "목공", 2, 1, 300000])
    ws_labor.append(["L-004", "철거", "일반", 2, 1, 200000])

    wb1.save(SAMPLES_DIR / "P001_최초견적.xlsx")

    # 2. P001 변경내역
    wb2 = openpyxl.Workbook()
    ws2 = wb2.active
    ws2.title = "변경견적"
    ws2.append(["SAMPLE — 데모용 가상 데이터"])
    ws2.append([])
    ws2.append(["항목코드", "구분", "항목명", "규격", "수량", "단위", "단가", "금액"])
    ws2.append(["E-001", "제작", "아크릴 케이스", "600x400x300 T5", 20, "EA", 92000, 1840000])
    ws2.append(["E-002", "제작", "목공 받침대", "자작나무 18T", 24, "EA", 120000, 2880000])
    ws2.append(["E-003", "인쇄", "그래픽 패널 출력", "UV 1200x900", 40, "EA", 45000, 1800000])
    ws2.append(["", "전기전자", "LED 조명모듈(고휘도)", "12V 3000K", 40, "EA", 22000, 880000])
    ws2.append(["E-005", "전기전자", "터치 센서 키트", "정전식", 20, "SET", 65000, 1300000])
    ws2.append(["E-006", "운송", "운송비", "5톤 윙바디", 3, "회", 450000, 1350000])
    ws2.append(["E-007", "디자인", "원형 디자인", "컨셉~CMF", 5, "일", 400000, 2000000])
    ws2.append(["E-008", "소프트웨어", "다국어 오디오 가이드", "4개 언어", 1, "SET", 1500000, 1500000])

    ws_appr = wb2.create_sheet(title="변경승인")
    ws_appr.append(["SAMPLE — 데모용 가상 데이터"])
    ws_appr.append([])
    ws_appr.append(["승인번호", "항목코드", "항목명", "변경내용", "승인여부", "승인일", "승인자"])
    ws_appr.append(["CR-01", "E-001", "아크릴 케이스", "소재 등급 상향", "승인", "2026-09-02", "김OO"])
    ws_appr.append(["CR-02", "E-006", "운송비", "운송 1회 추가", "반려", "2026-09-05", "김OO"])
    ws_appr.append(["CR-03", "E-008", "다국어 오디오 가이드", "신규 추가", "승인", "2026-09-10", "이OO"])
    ws_appr.append(["CR-04", "E-010", "안내 리플릿 인쇄", "항목 삭제", "승인", "2026-09-10", "이OO"])
    ws_appr.append(["CR-05", "E-002", "목공 받침대", "수량 추가", "", "", ""])
    ws_appr.append(["CR-06", "L-001", "설치 목공", "인원 1명 추가", "승인", "2026-09-12", "김OO"])
    ws_appr.append(["CR-07", "E-099", "포토존 조형물", "신규 검토", "승인", "2026-09-15", "이OO"])

    ws_chg_labor = wb2.create_sheet(title="변경인력")
    ws_chg_labor.append(["SAMPLE — 데모용 가상 데이터"])
    ws_chg_labor.append([])
    ws_chg_labor.append(["항목코드", "공정", "역할", "인원", "일수", "일단가"])
    ws_chg_labor.append(["L-001", "설치", "목공", 3, 3, 300000])
    ws_chg_labor.append(["L-002", "설치", "전기", 1, 3, 320000])
    ws_chg_labor.append(["L-003", "철거", "목공", 2, 1, 300000])
    ws_chg_labor.append(["L-005", "철거", "전기", 1, 1, 320000])

    wb2.save(SAMPLES_DIR / "P001_변경내역.xlsx")

def create_p002():
    # P002 (북미 USD)
    wb1 = openpyxl.Workbook()
    ws1 = wb1.active
    ws1.title = "견적"
    ws1.append(["SAMPLE — 데모용 가상 데이터"])
    ws1.append([])
    ws1.append(["항목코드", "구분", "항목명", "규격", "수량", "단위", "단가", "금액"])
    ws1.append(["U-01", "Hardware", "Acrylic Display", "Standard", 10, "EA", "$1,200", "$12,000"])
    ws1.append(["U-02", "Hardware", "Sensor Kit", "Standard", 5, "SET", "$450", "$2,250"])
    ws1.append(["U-03", "Content", "Translation", "English", 1, "LOT", "$3,000", "$3,000"])
    ws1.append(["U-04", "Service", "On-site Setup", "Days", 3, "DAYS", "TBD", "TBD"])
    wb1.save(SAMPLES_DIR / "P002_최초견적.xlsx")

    wb2 = openpyxl.Workbook()
    ws2 = wb2.active
    ws2.title = "변경견적"
    ws2.append(["SAMPLE — 데모용 가상 데이터"])
    ws2.append([])
    ws2.append(["항목코드", "구분", "항목명", "규격", "수량", "단위", "단가", "금액"])
    ws2.append(["U-01", "Hardware", "Acrylic Display", "Standard", 10, "EA", "$1,350", "$13,500"])
    ws2.append(["U-02", "Hardware", "Sensor Kit", "Standard", 8, "SET", "$450", "$3,600"])
    ws2.append(["U-03", "Content", "Translation", "English", 1, "LOT", "$3,000", "$3,000"])
    ws2.append(["U-04", "Service", "On-site Setup", "Days", 3, "DAYS", "$1,500", "$4,500"])

    ws_appr = wb2.create_sheet(title="변경승인")
    ws_appr.append(["SAMPLE — 데모용 가상 데이터"])
    ws_appr.append([])
    ws_appr.append(["승인번호", "항목코드", "항목명", "변경내용", "승인여부", "승인일", "승인자"])
    ws_appr.append(["CR-US-01", "U-01", "Acrylic Display", "Price update", "승인", "2026-09-01", "John D."])
    ws_appr.append(["CR-US-02", "U-02", "Sensor Kit", "Qty increase", "승인", "2026-09-02", "John D."])
    wb2.save(SAMPLES_DIR / "P002_변경내역.xlsx")

def create_broken_samples():
    wb1 = openpyxl.Workbook()
    ws1 = wb1.active
    ws1.title = "견적"
    ws1.append(["SAMPLE — 데모용 가상 데이터"])
    ws1.append([])
    ws1.append(["항목코드", "구분", "항목명", "수량"])
    ws1.append(["E-001", "제작", "테스트 항목", 10])
    wb1.save(BROKEN_DIR / "missing_columns.xlsx")

    wb2 = openpyxl.Workbook()
    ws2 = wb2.active
    ws2.title = "시트1"
    ws2.append(["임의의 데이터"])
    ws2_extra = wb2.create_sheet(title="요약")
    ws2_extra.append(["요약 정보"])
    wb2.save(BROKEN_DIR / "wrong_sheet_name.xlsx")

    fake_excel_path = BROKEN_DIR / "fake_excel.xlsx"
    with open(fake_excel_path, "w", encoding="utf-8") as f:
        f.write("이 파일은 실제 엑셀 파일이 아닌 일반 텍스트 파일입니다.")

def main():
    SAMPLES_DIR.mkdir(parents=True, exist_ok=True)
    BROKEN_DIR.mkdir(parents=True, exist_ok=True)
    create_p001()
    create_p002()
    create_broken_samples()
    print("가상 샘플 데이터 생성 완료!")

if __name__ == "__main__":
    main()
