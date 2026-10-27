# -*- coding: utf-8 -*-
import sys
from pathlib import Path
sys.path.append(str(Path(__file__).resolve().parent.parent))

import openpyxl
from auditor.loader import read_estimate_sheet
from auditor.categorizer import decompose_estimate_by_category, compare_decomposed_estimates

p_base = Path("data/samples/기준견적서4_전시_광고주컨펌.xlsx")
p_rev = Path("data/samples/전시_홍대점_변경내역.xlsx")
wb_base = openpyxl.load_workbook(p_base, data_only=True)
wb_rev = openpyxl.load_workbook(p_rev, data_only=True)

sheet_b = [s for s in wb_base.sheetnames if "수원" in s][0]
sheet_r = wb_rev.sheetnames[0]

items_base, _, _, _ = read_estimate_sheet(wb_base[sheet_b], p_base.name, sheet_b)
items_rev, _, _, _ = read_estimate_sheet(wb_rev[sheet_r], p_rev.name, sheet_r)

base_decomp = decompose_estimate_by_category(items_base)
rev_decomp = decompose_estimate_by_category(items_rev)

print(f"Base Total: KRW {base_decomp['total_amount']:,.0f} ({base_decomp['item_count']} items)")
print(f"Rev Total:  KRW {rev_decomp['total_amount']:,.0f} ({rev_decomp['item_count']} items)")

comp = compare_decomposed_estimates(base_decomp, rev_decomp)
print("\n=== 7대 공종별 분해 대조 결과 ===")
for c in comp:
    print(f"[{c['공종명']}]")
    print(f"  Base: KRW {c['기준금액']:>12,.0f} ({c['기준비중']:>5.1f}%) | Rev: KRW {c['신규금액']:>12,.0f} ({c['신규비중']:>5.1f}%)")
    print(f"  Diff: KRW {c['비용증감액']:>+12,.0f} ({c['증감률']:>+5.1f}%) | Status: {c['판정']}")
