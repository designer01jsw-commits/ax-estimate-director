# -*- coding: utf-8 -*-
import sys
from pathlib import Path
sys.path.append(str(Path(__file__).resolve().parent.parent))

import openpyxl
from auditor.projects import load_projects
from auditor.loader import (
    load_workbook_safe,
    find_best_estimate_sheet,
    read_estimate_sheet,
    read_labor_sheet,
    read_approval_sheet
)
from auditor.matching import match_items
from auditor.analysis import (
    analyze_estimates,
    analyze_labor,
    build_review_items,
    calculate_summary
)
from auditor.presentation_modules import (
    explain_cost_difference,
    get_traffic_light_status
)
from auditor.diagnostics import (
    normalize_baseline,
    diagnose_estimate_item,
    diagnose_labor_item,
    generate_executive_supervisor_opinion
)
from auditor.categorizer import (
    decompose_estimate_by_category,
    compare_decomposed_estimates,
    audit_all_new_items,
    generate_cross_project_audit_report
)

projects, _ = load_projects()
root_dir = Path(".")

# Test with project 3 (기준견적서4)
selected_proj = projects[3]
print("Testing with baseline project:", selected_proj["name"])

sample_est_str = selected_proj.get("sample_estimate")
candidate_orig = [
    root_dir / "data" / sample_est_str,
    root_dir / "data" / "samples" / sample_est_str,
    root_dir / "data" / "samples" / "전시 광고주컨펌완료 견적서 (1).xlsx"
]
orig_path = next((p for p in candidate_orig if p and p.exists()), None)
print("Resolved orig_path:", orig_path)

# Test multiple candidate comparison files
test_rev_files = [
    root_dir / "data" / "samples" / "전시_홍대점_변경내역.xlsx",
    root_dir / "data" / "samples" / "P001_변경내역.xlsx",
    root_dir / "data" / "samples" / "P002_변경내역.xlsx",
    root_dir / "data" / "samples" / "전시 광고주컨펌완료 견적서 (1).xlsx",
    root_dir / "data" / "samples" / "기준견적서2_키트컨텐츠_광고주컨펌.xlsx"
]

for rev_path in test_rev_files:
    if not rev_path.exists():
        continue
    print(f"\n==========================================")
    print(f"Testing upload of: {rev_path.name}")
    try:
        wb_orig = load_workbook_safe(orig_path, orig_path.name)
        wb_rev = load_workbook_safe(rev_path, rev_path.name)
        
        orig_sheets = list(wb_orig.sheetnames)
        rev_sheets = list(wb_rev.sheetnames)
        common_sheets = [s for s in orig_sheets if s in rev_sheets]
        
        if common_sheets:
            sheet_orig_selected = common_sheets[0]
            sheet_rev_selected = common_sheets[0]
        else:
            best_orig, _ = find_best_estimate_sheet(wb_orig, selected_proj.get("default_sheet"))
            best_rev, _ = find_best_estimate_sheet(wb_rev)
            sheet_orig_selected = best_orig
            sheet_rev_selected = best_rev
            
        print(f"Selected sheets: {sheet_orig_selected} vs {sheet_rev_selected}")
        
        ws_orig = wb_orig[sheet_orig_selected]
        ws_rev = wb_rev[sheet_rev_selected]
        
        orig_items, mismatches_orig, orig_dup, orig_warn = read_estimate_sheet(ws_orig, orig_path.name, sheet_orig_selected)
        rev_items, mismatches_rev, rev_dup, rev_warn = read_estimate_sheet(ws_rev, rev_path.name, sheet_rev_selected)
        
        print(f"Parsed items: orig={len(orig_items)}, rev={len(rev_items)}")
        
        # Labor sheet
        orig_labor, orig_lab_warn = [], []
        rev_labor, rev_lab_warn = [], []
        labor_candidates_orig = [s for s in orig_sheets if any(k in s for k in ["인력", "투입률", "공통"])]
        labor_candidates_rev = [s for s in rev_sheets if any(k in s for k in ["인력", "투입률", "공통"])]
        if labor_candidates_orig and labor_candidates_rev:
            try:
                orig_labor, _, orig_lab_warn = read_labor_sheet(wb_orig[labor_candidates_orig[0]], orig_path.name, labor_candidates_orig[0])
                rev_labor, _, rev_lab_warn = read_labor_sheet(wb_rev[labor_candidates_rev[0]], rev_path.name, labor_candidates_rev[0])
            except Exception as e:
                print("Labor parse failed (caught):", e)
                
        # Approval
        approval_records = []
        if "변경승인" in wb_rev.sheetnames:
            try:
                approval_records, _ = read_approval_sheet(wb_rev["변경승인"], rev_path.name, "변경승인")
            except Exception:
                pass
                
        matches = match_items(orig_items, rev_items, orig_dup, rev_dup, {})
        labor_matches = match_items(orig_labor, rev_labor) if (orig_labor or rev_labor) else []
        
        analyzed_estimates, orphan_approvals = analyze_estimates(matches, orig_items, rev_items, approval_records)
        analyzed_labor, used_lab_apprs = analyze_labor(labor_matches, approval_records)
        orphan_approvals = [a for a in orphan_approvals if a["row_key"] not in used_lab_apprs]

        all_warnings = orig_warn + rev_warn + orig_lab_warn + rev_lab_warn
        formula_mismatches = mismatches_orig + mismatches_rev
        review_items = build_review_items(
            analyzed_estimates, analyzed_labor, orphan_approvals, formula_mismatches, all_warnings
        )
        summary = calculate_summary(
            orig_items, rev_items, analyzed_estimates, analyzed_labor, review_items
        )
        
        for it in analyzed_estimates:
            tl_status, tl_desc = get_traffic_light_status(it)
            it["신호등"] = tl_status
            it["판정근거"] = tl_desc
            it["비싼이유_설명"] = explain_cost_difference(it, "KRW")
            
        baseline_norms = normalize_baseline(orig_items, orig_labor)
        diagnosed_estimates = [diagnose_estimate_item(it, baseline_norms) for it in analyzed_estimates]
        diagnosed_labor = [diagnose_labor_item(it, baseline_norms) for it in analyzed_labor]
        
        supervisor_opinion = generate_executive_supervisor_opinion(
            diagnosed_estimates=diagnosed_estimates,
            diagnosed_labor=diagnosed_labor,
            summary=summary,
            project_name=selected_proj.get("name", "프로젝트"),
            baseline_norms=baseline_norms
        )
        
        base_decomp = decompose_estimate_by_category(orig_items)
        rev_decomp = decompose_estimate_by_category(rev_items)
        category_comparison = compare_decomposed_estimates(base_decomp, rev_decomp)
        audited_new_items, total_excess_cost, fair_count, caution_count, excess_count = audit_all_new_items(
            rev_items, orig_items, base_decomp, currency="KRW"
        )
        cross_audit_opinion = generate_cross_project_audit_report(
            selected_proj.get("name", "기준 프로젝트"), orig_path.name, rev_path.name,
            base_decomp, rev_decomp, category_comparison,
            total_excess_cost, fair_count, caution_count, excess_count, currency="KRW"
        )
        print("SUCCESS! All logic executed without error for:", rev_path.name)
    except Exception as e:
        import traceback
        print(f"FAILED on {rev_path.name} with error:")
        traceback.print_exc()
