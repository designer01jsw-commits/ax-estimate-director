# -*- coding: utf-8 -*-
import sys
from pathlib import Path
sys.path.append(str(Path(__file__).resolve().parent.parent))

import openpyxl
from auditor.loader import read_estimate_sheet
from auditor.matching import match_items
from auditor.analysis import analyze_estimates

p_base = Path("data/samples/기준견적서4_전시_광고주컨펌.xlsx")
p_rev = Path("data/samples/전시_홍대점_변경내역.xlsx")
wb_base = openpyxl.load_workbook(p_base, data_only=True)
wb_rev = openpyxl.load_workbook(p_rev, data_only=True)

# Find sheet with 수원 in base
sheet_b = [s for s in wb_base.sheetnames if "수원" in s][0]
sheet_r = wb_rev.sheetnames[0]

items_base, _, d_b, _ = read_estimate_sheet(wb_base[sheet_b], p_base.name, sheet_b)
items_rev, _, d_r, _ = read_estimate_sheet(wb_rev[sheet_r], p_rev.name, sheet_r)

matches = match_items(items_base, items_rev, d_b, d_r)
analyzed, orphan = analyze_estimates(matches, items_base, items_rev, [])

print(f"Items Base: {len(items_base)}, Items Rev: {len(items_rev)}")
print(f"Matched pairs: {len(matches)}")
added_count = len([a for a in analyzed if a.get("변경유형") == "추가"])
deleted_count = len([a for a in analyzed if a.get("변경유형") == "삭제"])
unchanged_count = len([a for a in analyzed if a.get("변경유형") == "변동없음"])
diff_count = len([a for a in analyzed if a.get("변경유형") == "단가/수량 변경"])

print(f"Added (신규 추가): {added_count}")
print(f"Deleted (기존 삭제): {deleted_count}")
print(f"Unchanged (변동없음): {unchanged_count}")
print(f"Modified (단가/수량 변경): {diff_count}")
if matches:
    print("Sample matched pair:", matches[0]["orig_item"]["항목명"], "vs", matches[0]["rev_item"]["항목명"])
