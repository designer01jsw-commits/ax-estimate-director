# -*- coding: utf-8 -*-
"""
대용량 원본 엑셀 파일(총 150MB)에서 불필요한 미디어/TIFF 이미지를 제외하고
모든 시트와 셀 데이터만을 추출하여 초경량(총 200KB 미만) 클라우드 배포용 엑셀로 압축/변환합니다.
"""

import os
import openpyxl
from pathlib import Path

DATA_DIR = Path(r"c:\Users\ws081\OneDrive\문서\테스트_안티그래비티\AX대회\data")
SAMPLES_DIR = DATA_DIR / "samples"

TARGET_FILES = [
    "기준견적서_키트컨텐츠_광고주컨펌_260511.xlsx",
    "기준견적서2_키트컨텐츠_광고주컨펌.xlsx",
    "기준견적서3_키트컨텐츠_광고주컨펌.xlsx",
    "기준견적서4_전시_광고주컨펌.xlsx",
    "전시 광고주컨펌완료 견적서 (1).xlsx"
]

def compact_file(file_name: str):
    file_path = DATA_DIR / file_name
    if not file_path.exists():
        print(f"Skipping (not found): {file_name}")
        return

    orig_size_mb = file_path.stat().st_size / (1024 * 1024)
    print(f"\nProcessing: {file_name} ({orig_size_mb:.2f} MB)")

    wb_in = openpyxl.load_workbook(str(file_path), data_only=True)
    wb_out = openpyxl.Workbook()
    wb_out.remove(wb_out.active)  # Remove default empty sheet

    for sheet_name in wb_in.sheetnames:
        ws_in = wb_in[sheet_name]
        ws_out = wb_out.create_sheet(title=sheet_name)
        for row in ws_in.iter_rows(values_only=True):
            if any(cell is not None for cell in row):
                ws_out.append(list(row))

    wb_in.close()

    # Save compact file
    temp_target = DATA_DIR / f"temp_{file_name}"
    wb_out.save(str(temp_target))
    wb_out.close()

    new_size_kb = temp_target.stat().st_size / 1024
    print(f" -> Compacted size: {new_size_kb:.1f} KB (reduced by {(1 - (new_size_kb/1024)/orig_size_mb)*100:.1f}%)")

    # Replace original with compact version
    temp_target.replace(file_path)

    # Also copy to samples directory
    sample_copy = SAMPLES_DIR / file_name
    sample_copy.write_bytes(file_path.read_bytes())
    print(f" -> Synced to samples/: {sample_copy.name}")

def main():
    SAMPLES_DIR.mkdir(parents=True, exist_ok=True)
    for fname in TARGET_FILES:
        compact_file(fname)
    print("\nAll target Excel files compacted successfully!")

if __name__ == "__main__":
    main()
