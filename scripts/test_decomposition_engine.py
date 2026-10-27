# -*- coding: utf-8 -*-
import sys
from pathlib import Path
sys.path.append(str(Path(__file__).resolve().parent.parent))

import openpyxl
from auditor.loader import load_workbook_safe, read_estimate_sheet, read_labor_sheet

# Test reading sheets
p_base = Path("data/samples/기준견적서4_전시_광고주컨펌.xlsx")
p_rev = Path("data/samples/전시_홍대점_변경내역.xlsx")

wb_base = load_workbook_safe(str(p_base), p_base.name)
wb_rev = load_workbook_safe(str(p_rev), p_rev.name)

print("Base sheets:", wb_base.sheetnames)
print("Rev sheets:", wb_rev.sheetnames)

# In base, let's try reading '1. 수원 사업장'
ws_base = wb_base["1. 수원 사업장"]
items_base, _, _, warn_base = read_estimate_sheet(ws_base, p_base.name, "1. 수원 사업장")
print(f"Base items count: {len(items_base)}")
if items_base:
    print("Sample base item:", items_base[0])

# In rev, try reading '견적'
ws_rev = wb_rev[wb_rev.sheetnames[0]]
items_rev, _, _, warn_rev = read_estimate_sheet(ws_rev, p_rev.name, wb_rev.sheetnames[0])
print(f"Rev items count: {len(items_rev)}")
if items_rev:
    print("Sample rev item:", items_rev[0])
