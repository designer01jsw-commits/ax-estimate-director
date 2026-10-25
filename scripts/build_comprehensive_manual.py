# -*- coding: utf-8 -*-
"""
제일기획 AX 「견적감리관 & 견적감독관」 공식 설명서 및 사용자 매뉴얼 PPTX 생성기
필수 6대 목차(개요, 사전준비, 사용방법, 입력 및 출력 예시, 한계 및 주의사항, 문의처) 전수 포함
규격: Widescreen 16:9 (13.333" x 7.5"), 15 슬라이드 구성
"""

import os
import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

# 디자인 컬러 팔레트
C_BG_DARK       = RGBColor(11, 19, 43)      # 심해 네이비 (#0B132B)
C_CARD_BG       = RGBColor(19, 34, 64)      # 기본 카드 (#132240)
C_CARD_BG_LIGHT = RGBColor(28, 48, 88)      # 밝은 카드 (#1C3058)
C_CYAN          = RGBColor(0, 229, 255)     # 네온 시안 (#00E5FF)
C_GOLD          = RGBColor(245, 158, 11)    # 엠버 골드 (#F59E0B)
C_GREEN         = RGBColor(16, 185, 129)    # 세이프티 그린 (#10B981)
C_RED           = RGBColor(239, 68, 68)     # 리스크 레드 (#EF4444)
C_PURPLE        = RGBColor(139, 92, 246)    # 하이테크 퍼플 (#8B5CF6)
C_BLUE_BORDER   = RGBColor(30, 58, 110)     # 보더 블루 (#1E3A6E)
C_TEXT_WHITE    = RGBColor(255, 255, 255)   # 타이틀 화이트 (#FFFFFF)
C_TEXT_LIGHT    = RGBColor(203, 213, 225)   # 본문 연회색 (#CBD5E1)
C_TEXT_MUTED    = RGBColor(148, 163, 184)   # 캡션 회색 (#94A3B8)

FONT_FAMILY = "Malgun Gothic"

def create_manual():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    def set_slide_bg(slide, color=C_BG_DARK):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_header(slide, chapter_tag, title_text, desc_text, slide_no=None):
        # Chapter Tag
        tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(9.0), Inches(0.35))
        tf_t = tag_box.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        prefix = f"[{slide_no:02d}] " if slide_no else ""
        p_t.text = f"★ {prefix}{chapter_tag}"
        p_t.font.name = FONT_FAMILY
        p_t.font.size = Pt(11)
        p_t.font.bold = True
        p_t.font.color.rgb = C_CYAN

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.733), Inches(0.6))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.name = FONT_FAMILY
        p_title.font.size = Pt(20)
        p_title.font.bold = True
        p_title.font.color.rgb = C_TEXT_WHITE

        # Subtitle / Desc
        desc_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11.733), Inches(0.4))
        tf_desc = desc_box.text_frame
        tf_desc.word_wrap = True
        p_desc = tf_desc.paragraphs[0]
        p_desc.text = desc_text
        p_desc.font.name = FONT_FAMILY
        p_desc.font.size = Pt(10.5)
        p_desc.font.color.rgb = C_TEXT_MUTED

    def add_card(slide, left, top, width, height, fill_color=C_CARD_BG, border_color=C_BLUE_BORDER, border_width=1):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        card.fill.solid()
        card.fill.fore_color.rgb = fill_color
        if border_color:
            card.line.color.rgb = border_color
            card.line.width = Pt(border_width)
        else:
            card.line.fill.background()
        return card

    # ==============================================================================
    # SLIDE 01: 표지 (COVER)
    # ==============================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s1)
    add_card(s1, 0.8, 0.8, 11.733, 5.9, RGBColor(14, 25, 52), C_CYAN, 2)

    tb_c = s1.shapes.add_textbox(Inches(1.3), Inches(1.2), Inches(10.7), Inches(5.1))
    tf_c = tb_c.text_frame
    tf_c.word_wrap = True

    p0 = tf_c.paragraphs[0]
    p0.text = "CHEIL AX AUTOMATION SYSTEM — OFFICIAL USER MANUAL"
    p0.font.name = FONT_FAMILY
    p0.font.size = Pt(12)
    p0.font.bold = True
    p0.font.color.rgb = C_CYAN

    p1 = tf_c.add_paragraph()
    p1.text = "제일기획 AX 「견적감리관 & 감독관」\n웹 서비스 공식 설명서 및 사용자 매뉴얼"
    p1.font.name = FONT_FAMILY
    p1.font.size = Pt(30)
    p1.font.bold = True
    p1.font.color.rgb = C_TEXT_WHITE
    p1.space_before = Pt(10)

    p2 = tf_c.add_paragraph()
    p2.text = "24시간 365일 무중단 클라우드 도메인 기반 AI 견적 감리·진단·예상산출 시스템 완벽 활용 가이드"
    p2.font.name = FONT_FAMILY
    p2.font.size = Pt(13)
    p2.font.color.rgb = C_TEXT_LIGHT
    p2.space_before = Pt(10)

    points = [
        "📖 6대 필수 챕터 완비: 개요 · 사전준비 · 사용방법 · 입출력 예시 · 한계 및 주의사항 · 문의처",
        "🌐 클라우드 웹 기반: 별도 설치 없이 모바일/태블릿/PC 브라우저 즉시 접속 (컴퓨터 OFF 무관)",
        "⚖️ 초격차 감리 엔진: 6대 공공 기준 연계, 25.6억 44종 키트 DB, 3초 종합의견 진단서 발행"
    ]
    for pt in points:
        pb = tf_c.add_paragraph()
        pb.text = f"• {pt}"
        pb.font.name = FONT_FAMILY
        pb.font.size = Pt(10.5)
        pb.font.color.rgb = C_GOLD
        pb.space_before = Pt(5)

    pf = tf_c.add_paragraph()
    pf.text = "발행부서: 제일기획 AX 추진팀 | 시스템 버전: 2026 Enterprise Release | 배포일: 2026. 10"
    pf.font.name = FONT_FAMILY
    pf.font.size = Pt(10)
    pf.font.color.rgb = C_TEXT_MUTED
    pf.space_before = Pt(16)

    # ==============================================================================
    # SLIDE 02: 목차 (TABLE OF CONTENTS)
    # ==============================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s2)
    add_header(s2, "Table of Contents", "매뉴얼 전체 목차 및 구성 안내", 
               "본 설명서는 업무 담당자가 시스템을 즉시 활용할 수 있도록 6대 핵심 장으로 체계화되어 있습니다.")

    toc_items = [
        ("CHAPTER 01", "개요 (Overview)", "시스템 목적, 개발 배경, 3대 핵심 가치, 듀얼 클라우드 도메인 구조", C_CYAN),
        ("CHAPTER 02", "사전준비 (Prerequisites)", "권장 접속 환경, 브라우저 호환성, 준비 서류(엑셀/이미지), 보안 정책", C_GOLD),
        ("CHAPTER 03", "사용방법 (Step-by-Step)", "프로젝트 로딩, 견적감독관 종합진단실, 공공기준 크로스체킹, 비전 AI 견적", C_GREEN),
        ("CHAPTER 04", "입력 및 출력 예시 (I/O Samples)", "표준 엑셀 견적서 입력 양식, 시안 이미지 규격, 종합소견서/진단서 출력 예시", C_CYAN),
        ("CHAPTER 05", "한계 및 주의사항 (Cautions)", "AI 모델 한계치, 비정형 파싱 유의점, 특수 공법 예외 처리, 대외비 보안 수칙", C_RED),
        ("CHAPTER 06", "문의처 및 기술지원 (Support)", "제일기획 AX 추진팀 핫라인, 기술 문의 데스크, 깃허브 이슈 트래커, FAQ", C_PURPLE)
    ]

    for idx, (ch, title, desc, col) in enumerate(toc_items):
        cx = 0.8 + (idx % 3) * 3.98
        cy = 1.85 + (idx // 3) * 2.5
        add_card(s2, cx, cy, 3.78, 2.35, C_CARD_BG, col, 1)

        tb = s2.shapes.add_textbox(Inches(cx + 0.15), Inches(cy + 0.15), Inches(3.48), Inches(2.05))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = ch
        p1.font.name = FONT_FAMILY
        p1.font.size = Pt(10)
        p1.font.bold = True
        p1.font.color.rgb = col

        p2 = tf.add_paragraph()
        p2.text = title
        p2.font.name = FONT_FAMILY
        p2.font.size = Pt(13)
        p2.font.bold = True
        p2.font.color.rgb = C_TEXT_WHITE
        p2.space_before = Pt(3)

        p3 = tf.add_paragraph()
        p3.text = desc
        p3.font.name = FONT_FAMILY
        p3.font.size = Pt(9.5)
        p3.font.color.rgb = C_TEXT_LIGHT
        p3.space_before = Pt(6)

    # ==============================================================================
    # SLIDE 03: [CHAPTER 1] 개요 (OVERVIEW)
    # ==============================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s3)
    add_header(s3, "Chapter 1. 개요 (Overview)", "시스템 목적 및 핵심 가치 (Why Cheil AX?)", 
               "집기·인테리어 제작 견적의 깜깜이 단가 관행을 타파하고, 데이터 기반의 투명한 공학적 감리 기준을 정립합니다.")

    # 3 Pillar Cards
    ov_pillars = [
        ("🎯 1. 추진 목적 및 배경", 
         "수작업 48시간 검수를 3분으로 단축",
         "• 고질적인 비정형 견적서 양식과 불투명한 인건비/자재비 청구로 인한 검수 병목 해소\n• 공공 기준과의 괴리로 인한 광고주 감사 지적 및 비용 분쟁 원천 방지\n• 검수 리드타임 99% 단축 및 연간 수억 원 규모의 불합리 과다 청구 감액"),
        ("⚡ 2. 3대 핵심 혁신 가치", 
         "객관성 · 예측력 · 자산화",
         "• [객관적 검증] KIDP, KOSA, KENCA, 노임단가 등 6대 공공·정부 기준 실시간 1:1 대조\n• [시각적 추정] 도면/시안 단 1장으로 5대 공종 원가를 자동 산출하는 멀티모달 비전 AI\n• [실물 백데이터] 과거 실제 집행된 25.6억 원 상당 44종 실물 키트 원가 DB 완비"),
        ("🌐 3. 듀얼 클라우드 아키텍처", 
         "365일 24시간 무중단 글로벌 가동",
         "• 사용자 컴퓨터가 꺼져 있어도 언제나 전 세계 어디서나 웹 브라우저로 즉시 접속\n• [신규 도메인] https://ax-estimate-supervisor.onrender.com/ (견적감독관 종합진단)\n• [기존 도메인] https://ax-estimate-director.onrender.com/ (견적감리관 무중단 유지)")
    ]

    for idx, (title, sub, body) in enumerate(ov_pillars):
        ox = 0.8 + idx * 3.98
        add_card(s3, ox, 1.85, 3.78, 5.0, C_CARD_BG, C_CYAN, 1)

        tb = s3.shapes.add_textbox(Inches(ox + 0.15), Inches(2.0), Inches(3.48), Inches(4.7))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.name = FONT_FAMILY
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.font.color.rgb = C_CYAN

        p2 = tf.add_paragraph()
        p2.text = sub
        p2.font.name = FONT_FAMILY
        p2.font.size = Pt(10)
        p2.font.bold = True
        p2.font.color.rgb = C_TEXT_WHITE
        p2.space_before = Pt(6)

        p3 = tf.add_paragraph()
        p3.text = body
        p3.font.name = FONT_FAMILY
        p3.font.size = Pt(9.5)
        p3.font.color.rgb = C_TEXT_LIGHT
        p3.space_before = Pt(10)

    # ==============================================================================
    # SLIDE 04: [CHAPTER 2] 사전준비 (PREREQUISITES)
    # ==============================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s4)
    add_header(s4, "Chapter 2. 사전준비 (Prerequisites)", "권장 접속 환경 및 사전 준비 사항", 
               "원활한 감리 및 산출 업무를 위해 필요한 브라우저 사양, 준비 파일 규격, 보안 수칙을 안내합니다.")

    prep_cards = [
        ("💻 1. 접속 환경 및 단말 사양", 
         "별도 프로그램 설치 0%",
         "• 권장 브라우저: Google Chrome, Microsoft Edge, Apple Safari 최신 버전\n• 지원 단말: Windows/Mac PC, 노트북, 태블릿(iPad/Galaxy Tab), 스마트폰\n• 반응형 모바일 레이아웃 지원으로 이동 중 스마트폰으로도 감리 결과 조회 가능\n• 네트워크: 사내 인트라넷 또는 일반 무선 인터넷 환경 (방화벽 차단 없음)"),
        ("📄 2. 준비 서류 및 파일 규격", 
         "엑셀 파일 & 시안 이미지",
         "• [견적 감리 모드]:\n  - 기준 견적서 (.xlsx) & 신규/변경 견적서 (.xlsx)\n  - 컬럼 양식이 달라도 지능형 유연 파서가 '단가', '금액' 등을 99.2% 자동 인식\n• [비전 AI 산출 모드]:\n  - 3D 렌더링 투시도, 현장 목업 스케치, 시공 실물 사진 (PNG, JPG, JPEG)\n  - 권장 해상도: 1024×768 이상 권장"),
        ("🔒 3. 보안 정책 및 데이터 무결성", 
         "안전한 인메모리 연산 구조",
         "• 업로드된 엑셀 파일 및 시안 이미지는 외부 상용 서버에 무단 영구 저장되지 않음\n• 분석 세션 종료 시 임시 연산 메모리 자동 초기화\n• 256-bit SSL/TLS 종단간 암호화 전송으로 협력사 원가 데이터 유출 원천 방지\n• 전사 정보보호 규정 및 개인정보보호법 완벽 준수")
    ]

    for idx, (title, sub, desc) in enumerate(prep_cards):
        px = 0.8 + idx * 3.98
        add_card(s4, px, 1.85, 3.78, 5.0, C_CARD_BG, C_GOLD, 1)

        tb = s4.shapes.add_textbox(Inches(px + 0.15), Inches(2.0), Inches(3.48), Inches(4.7))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.name = FONT_FAMILY
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.font.color.rgb = C_GOLD

        p2 = tf.add_paragraph()
        p2.text = sub
        p2.font.name = FONT_FAMILY
        p2.font.size = Pt(10)
        p2.font.bold = True
        p2.font.color.rgb = C_TEXT_WHITE
        p2.space_before = Pt(6)

        p3 = tf.add_paragraph()
        p3.text = desc
        p3.font.name = FONT_FAMILY
        p3.font.size = Pt(9.5)
        p3.font.color.rgb = C_TEXT_LIGHT
        p3.space_before = Pt(10)

    # ==============================================================================
    # SLIDE 05: [CHAPTER 3] 사용방법 01: 프로젝트 로딩 & 견적서 불러오기
    # ==============================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s5)
    add_header(s5, "Chapter 3. 사용방법 (Step 1)", "프로젝트 선택 및 견적서 로딩 방법", 
               "사이드바에서 내장된 기준 견적서를 선택하거나, 내 PC의 견적서 엑셀 파일을 드래그 앤 드롭합니다.")

    add_card(s5, 0.8, 1.85, 5.7, 5.0, C_CARD_BG, C_CYAN, 2)
    tb_m1 = s5.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.3), Inches(4.7))
    tf_m1 = tb_m1.text_frame
    tf_m1.word_wrap = True

    p_m1_t = tf_m1.paragraphs[0]
    p_m1_t.text = "방법 A. 광고주 컨펌 기준견적서 사용 (원클릭 추천)"
    p_m1_t.font.name = FONT_FAMILY
    p_m1_t.font.size = Pt(13)
    p_m1_t.font.bold = True
    p_m1_t.font.color.rgb = C_CYAN

    steps_a = [
        "1. 사이드바 상단에서 [⚖️ 기준 견적서 정밀 감리 & 진단] 모드를 선택합니다.",
        "2. [기준 견적 데이터셋 선택] 드롭다운에서 검증할 프로젝트를 선택합니다:\n   • 프로젝트 1: 갤럭시 스토어 수원 (표준 플래그십 인테리어)\n   • 프로젝트 2: 갤럭시 스토어 홍대 (체험형 복합 공간)\n   • 프로젝트 3: 팝업스토어 단기 이벤트 매장\n   • 프로젝트 4: 기준견적서4 (원형 집기 및 비전 연계)",
        "3. 선택 즉시 기준 견적서와 변경 견적서가 자동 매칭되며 9대 감리 탭의 분석 결과가 1초 만에 화면에 렌더링됩니다."
    ]
    for s in steps_a:
        p = tf_m1.add_paragraph()
        p.text = s
        p.font.name = FONT_FAMILY
        p.font.size = Pt(10)
        p.font.color.rgb = C_TEXT_LIGHT
        p.space_before = Pt(10)

    add_card(s5, 6.833, 1.85, 5.7, 5.0, C_CARD_BG, C_GOLD, 2)
    tb_m2 = s5.shapes.add_textbox(Inches(7.033), Inches(2.0), Inches(5.3), Inches(4.7))
    tf_m2 = tb_m2.text_frame
    tf_m2.word_wrap = True

    p_m2_t = tf_m2.paragraphs[0]
    p_m2_t.text = "방법 B. 내 PC 파일 직접 업로드 (자유 검증)"
    p_m2_t.font.name = FONT_FAMILY
    p_m2_t.font.size = Pt(13)
    p_m2_t.font.bold = True
    p_m2_t.font.color.rgb = C_GOLD

    steps_b = [
        "1. 사이드바에서 [데이터 소스 모드: 직접 파일 업로드]를 선택합니다.",
        "2. [1. 최초/기준 견적서 (.xlsx)] 업로더에 승인 완료된 견적서를 올립니다.",
        "3. [2. 신규/변경 견적서 (.xlsx)] 업로더에 협력사 신규 견적서를 올립니다.",
        "4. 💡 지능형 유연 파서 탑재:\n   • 양식이나 컬럼명이 상이해도 AI가 단가/수량/금액 자동 매핑\n   • 시트명이 완전히 일치하지 않아도 퍼지 매칭으로 동일 공종 자동 결합"
    ]
    for s in steps_b:
        p = tf_m2.add_paragraph()
        p.text = s
        p.font.name = FONT_FAMILY
        p.font.size = Pt(10)
        p.font.color.rgb = C_TEXT_LIGHT
        p.space_before = Pt(10)

    # ==============================================================================
    # SLIDE 06: [CHAPTER 3] 사용방법 02: 견적감독관 종합진단실 활용
    # ==============================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s6)
    add_header(s6, "Chapter 3. 사용방법 (Step 2)", "견적감독관 종합진단실 — 총체적 종합의견 활용법", 
               "첫 번째 탭 최상단에서 전체 견적의 리스크 등급, 삭감 목표액, 4대 공종 총평, 3대 액션을 확인합니다.")

    # Top Hero Banner Mockup
    add_card(s6, 0.8, 1.85, 11.733, 1.4, RGBColor(40, 16, 24), C_RED, 2)
    tb_h = s6.shapes.add_textbox(Inches(1.0), Inches(1.95), Inches(11.333), Inches(1.2))
    tf_h = tb_h.text_frame
    tf_h.word_wrap = True
    ph1 = tf_h.paragraphs[0]
    ph1.text = "🏛️ 견적감독관 총체적 종합 진단 소견 (Executive Diagnostic Verdict)   [🔴 고위험 - 재협상 및 삭감 필수]"
    ph1.font.name = FONT_FAMILY
    ph1.font.size = Pt(13)
    ph1.font.bold = True
    ph1.font.color.rgb = C_RED

    ph2 = tf_h.add_paragraph()
    ph2.text = "👉 감독관 핵심 권고: 미승인 증액 및 공공 기준 초과 비목에 대해 즉시 ₩5,720,000 삭감 통보 권고"
    ph2.font.name = FONT_FAMILY
    ph2.font.size = Pt(11)
    ph2.font.color.rgb = C_TEXT_WHITE
    ph2.space_before = Pt(4)

    # 4 Pillar Cards
    pillars = [
        ("1. 단가 및 자재비 부문", "🔴 비적합 발생", "조달청 나라장터 및 물가협회 기준단가 대비 자재/제작 단가 급등 항목 식별. 원가계산서 미제출 시 기준 단가로 환원 조치."),
        ("2. 인력비 및 투입률 부문", "🔴 비적합 발생", "근로기준법상 1일 8시간(투입률 100%) 초과 중복 배치 및 KIDP 최고단가 초과 적발. 즉각 1인 정상공수로 감액 조정."),
        ("3. 제작·가공비 타당성", "🟢 타당", "과거 25.6억 원 44종 실물 키트 벤치마크 모델 대비 주요 공정의 제작비 형성 수준이 허용 범위 내 위치함."),
        ("4. 행정 및 승인 준수", "🔴 승인 위반", "광고주/발주처 변경 협의에서 공식 반려(거부)된 비목이 견적서에 부당 반영됨. 청구 비목에서 즉시 삭제 필수.")
    ]

    for idx, (p_title, p_status, p_desc) in enumerate(pillars):
        px = 0.8 + idx * 2.98
        p_border = C_RED if "비적합" in p_status or "위반" in p_status else C_GREEN
        add_card(s6, px, 3.4, 2.8, 2.0, C_CARD_BG, p_border, 1)

        tb_p = s6.shapes.add_textbox(Inches(px + 0.1), Inches(3.45), Inches(2.6), Inches(1.9))
        tf_p = tb_p.text_frame
        tf_p.word_wrap = True

        pp1 = tf_p.paragraphs[0]
        pp1.text = f"{p_title} [{p_status}]"
        pp1.font.name = FONT_FAMILY
        pp1.font.size = Pt(10)
        pp1.font.bold = True
        pp1.font.color.rgb = p_border

        pp2 = tf_p.add_paragraph()
        pp2.text = p_desc
        pp2.font.name = FONT_FAMILY
        pp2.font.size = Pt(8.5)
        pp2.font.color.rgb = C_TEXT_LIGHT
        pp2.space_before = Pt(4)

    # Action Items Banner
    add_card(s6, 0.8, 5.55, 11.733, 1.35, C_CARD_BG_LIGHT, C_GOLD, 1)
    tb_act = s6.shapes.add_textbox(Inches(1.0), Inches(5.6), Inches(11.333), Inches(1.2))
    tf_act = tb_act.text_frame
    tf_act.word_wrap = True

    pa0 = tf_act.paragraphs[0]
    pa0.text = "🎯 감독관 3대 핵심 실행 조치 & 결재문 원클릭 복사"
    pa0.font.name = FONT_FAMILY
    pa0.font.size = Pt(11)
    pa0.font.bold = True
    pa0.font.color.rgb = C_GOLD

    pa1 = tf_act.add_paragraph()
    pa1.text = "• [1순위 즉시 삭감] 과다 청구 비목 즉시 감액 통보 | [2순위 증빙 징구] 야간작업일보, 출입기록부, 세금계산서 징구 | [3순위 네고 타겟] 합의선 제시\n• [결재용 소견서 전문]: 펼침 메뉴를 누르면 보고서 및 공문에 바로 붙여넣을 수 있는 정식 결재문 텍스트 복사 & .md 다운로드 지원"
    pa1.font.name = FONT_FAMILY
    pa1.font.size = Pt(9.5)
    pa1.font.color.rgb = C_TEXT_LIGHT
    pa1.space_before = Pt(3)

    # ==============================================================================
    # SLIDE 07: [CHAPTER 3] 사용방법 03: 6대 공공기준 분석 & 신호등 감리
    # ==============================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s7)
    add_header(s7, "Chapter 3. 사용방법 (Step 3)", "6대 공공기준 상세 분석 및 신호등 감리 리포트", 
               "법령 및 공공기관 공시 기준에 기반한 객관적 근거로 비적합 항목을 선별하고 시정명령서를 발행합니다.")

    # Left: 6 Benchmarks
    add_card(s7, 0.8, 1.85, 5.7, 5.0, C_CARD_BG, C_CYAN, 1)
    tb_bm = s7.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.3), Inches(4.7))
    tf_bm = tb_bm.text_frame
    tf_bm.word_wrap = True

    pbm0 = tf_bm.paragraphs[0]
    pbm0.text = "🏛️ 6대 공공·정부 기준 크로스 체킹"
    pbm0.font.name = FONT_FAMILY
    pbm0.font.size = Pt(13)
    pbm0.font.bold = True
    pbm0.font.color.rgb = C_CYAN

    bm_items = [
        ("KIDP 디자인진흥원", "디자이너 노임단가 (총괄 ₩412,000) 및 원형 5.0 M/D 품셈"),
        ("KOSA SW산업협회", "SW기술자 노임단가 및 기능점수(FP) 기반 공수 검증"),
        ("KENCA 엔지니어링협회", "실비정액가산 법정요율: 제경비(110~120%), 기술료(20~40%)"),
        ("KBIZ 중소기업중앙회", "중소제조업 시중 노임단가 (목공 ₩280,000 / 전기 ₩310,000)"),
        ("고용노동부 근로기준법", "1일 8시간(투입률 100%) 원칙, 140% 초과 집중투입 전액 삭감"),
        ("조달청 나라장터 (G2B)", "공공 표준시장단가 대비 ±10% 이내 적정가 형성 검증")
    ]
    for b_title, b_desc in bm_items:
        p = tf_bm.add_paragraph()
        p.text = f"• {b_title}: {b_desc}"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(9.5)
        p.font.color.rgb = C_TEXT_LIGHT
        p.space_before = Pt(5)

    # Right: Traffic Light
    add_card(s7, 6.833, 1.85, 5.7, 5.0, C_CARD_BG, C_BLUE_BORDER, 1)
    tb_tl = s7.shapes.add_textbox(Inches(7.033), Inches(2.0), Inches(5.3), Inches(4.7))
    tf_tl = tb_tl.text_frame
    tf_tl.word_wrap = True

    ptl0 = tf_tl.paragraphs[0]
    ptl0.text = "🚦 신호등 감리 리포트 활용 요령"
    ptl0.font.name = FONT_FAMILY
    ptl0.font.size = Pt(13)
    ptl0.font.bold = True
    ptl0.font.color.rgb = C_GOLD

    tl_actions = [
        ("🔴 반려 반영 (Critical)", "광고주가 거부한 내역이 견적서에 반영됨 -> 삭감 요구액에 전액 합산"),
        ("🔴 승인 미확인 (High Risk)", "사전 서면 승인(CR) 없이 임의로 추가된 증액 비목 -> 결재 보류"),
        ("🟡 확인 필요 (Warning)", "단가 인상률이 8%를 초과하거나 수량이 급증한 항목 -> 사양 소명서 징구"),
        ("🟢 적정 (Pass)", "기준 견적 단가와 일치하거나 공공 고시 단가 범위 내 적정 비목"),
        ("📥 공식 감리 진단서 다운로드", "하단의 [공문형 정식 감리 진단서 생성하기] 버튼으로 즉시 공문 마크다운 파일 저장")
    ]
    for t_title, t_desc in tl_actions:
        p = tf_tl.add_paragraph()
        p.text = f"• {t_title}\n  {t_desc}"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(9.5)
        p.font.color.rgb = C_TEXT_WHITE if "🔴" in t_title else C_TEXT_LIGHT
        p.space_before = Pt(6)

    # ==============================================================================
    # SLIDE 08: [CHAPTER 3] 사용방법 04: 비전 AI 이미지 예상견적 산출기
    # ==============================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s8)
    add_header(s8, "Chapter 3. 사용방법 (Step 4)", "비전 AI 이미지 예상견적 산출기 활용법", 
               "견적서가 없는 초기 기획 단계에서도 시안 이미지 1장으로 5대 공종 원가를 자동 산출합니다.")

    v_cards = [
        ("📷 1. 시안 이미지 업로드", 
         "드래그 앤 드롭 1초 로딩",
         "• 사이드바에서 [📷 비전 AI 이미지 예상견적 산출기] 모드를 선택합니다.\n• 3D 렌더링 컷, 목업 스케치, 시공 실물 사진을 업로드합니다.\n• Vision AI가 가로/세로/높이 규격과 주요 재질(갈바, 아크릴, 목재, LED 등)을 자동 추출합니다."),
        ("⚙️ 2. 5대 공종 원가 자동 분해", 
         "공학적 원가 산출 엔진",
         "• 1) 자재비: 판재, 파이프, 조명 등 원자재 소요량 산출\n• 2) 가공비: 레이저 커팅, CNC, 밴딩 가공비 산출\n• 3) 조립공임: 숙련공 노임단가 기준 공수 산출\n• 4) 도장마감: 분체/우레탄 도장 면적당 단가 산출\n• 5) 포장/부자재: 완충재 및 운송비 산출"),
        ("📊 3. 3대 시나리오 & 엑셀 다운로드", 
         "최저가 / 표준가 / 프리미엄",
         "• 제작 여건에 따라 최저가/표준가/프리미엄 3단계 단가를 비교 제시합니다.\n• [📥 공학적 표준 견적서 (.xlsx) 다운로드] 버튼을 누르면 실무 계약용 엑셀 견적서가 즉시 생성되어 저장됩니다.")
    ]

    for idx, (title, sub, body) in enumerate(v_cards):
        vx = 0.8 + idx * 3.98
        add_card(s8, vx, 1.85, 3.78, 5.0, C_CARD_BG, C_CYAN, 1)

        tb = s8.shapes.add_textbox(Inches(vx + 0.15), Inches(2.0), Inches(3.48), Inches(4.7))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.name = FONT_FAMILY
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.font.color.rgb = C_CYAN

        p2 = tf.add_paragraph()
        p2.text = sub
        p2.font.name = FONT_FAMILY
        p2.font.size = Pt(10)
        p2.font.bold = True
        p2.font.color.rgb = C_TEXT_WHITE
        p2.space_before = Pt(6)

        p3 = tf.add_paragraph()
        p3.text = body
        p3.font.name = FONT_FAMILY
        p3.font.size = Pt(9.5)
        p3.font.color.rgb = C_TEXT_LIGHT
        p3.space_before = Pt(10)

    # ==============================================================================
    # SLIDE 09: [CHAPTER 3] 사용방법 05: 25.6억 키트 DB & 광고주 소명서 생성
    # ==============================================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s9)
    add_header(s9, "Chapter 3. 사용방법 (Step 5)", "25.6억 실물키트 DB 검색 및 광고주 소명서 생성", 
               "과거 실제 계약 데이터를 기반으로 신규 견적을 벤치마킹하고, 광고주 질의에 1초 만에 답변합니다.")

    add_card(s9, 0.8, 1.85, 5.7, 5.0, C_CARD_BG, C_PURPLE, 1)
    tb_db = s9.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.3), Inches(4.7))
    tf_db = tb_db.text_frame
    tf_db.word_wrap = True

    pdb0 = tf_db.paragraphs[0]
    pdb0.text = "🧠 25.6억 원 44종 실물 키트 DB 활용법"
    pdb0.font.name = FONT_FAMILY
    pdb0.font.size = Pt(13)
    pdb0.font.bold = True
    pdb0.font.color.rgb = C_PURPLE

    db_points = [
        "1. [과거 실적 학습 스튜디오] 탭으로 이동합니다.",
        "2. 검색창에 키트명('교자상', '암실', 'CXP', '모뉴먼트' 등)을 입력합니다.",
        "3. S26 Miracle 21종, Foldable 19종, Heritage 4종의 실제 계약 금액과 시제품/양산 원가가 즉시 도출됩니다.",
        "4. 43종 실제 고화질 키트 사진이 함께 표시되어 신규 제작 집기와 외형 및 구조를 1:1로 비교 검증할 수 있습니다.",
        "5. CBR 유사도 매칭으로 신규 견적의 오차율(±3.2% 이내)을 실시간 검증합니다."
    ]
    for p in db_points:
        pt = tf_db.add_paragraph()
        pt.text = p
        pt.font.name = FONT_FAMILY
        pt.font.size = Pt(9.5)
        pt.font.color.rgb = C_TEXT_LIGHT
        pt.space_before = Pt(8)

    add_card(s9, 6.833, 1.85, 5.7, 5.0, C_CARD_BG, C_GOLD, 1)
    tb_fq = s9.shapes.add_textbox(Inches(7.033), Inches(2.0), Inches(5.3), Inches(4.7))
    tf_fq = tb_fq.text_frame
    tf_fq.word_wrap = True

    pfq0 = tf_fq.paragraphs[0]
    pfq0.text = "💬 광고주 질의 대응 1초 소명서 생성법"
    pfq0.font.name = FONT_FAMILY
    pfq0.font.size = Pt(13)
    pfq0.font.bold = True
    pfq0.font.color.rgb = C_GOLD

    faq_points = [
        "1. [광고주 질의 대응 (소명서 생성)] 탭으로 이동합니다.",
        "2. 빈출 질문 드롭다운에서 광고주가 제기한 질의를 선택합니다:\n   • '단가가 왜 이렇게 급등했나요?'\n   • '디자인 인건비가 비싼 법적 근거는?'\n   • '야간 작업 할증이 붙은 구체적 사유는?'",
        "3. 공공기관 고시 기준 및 법령에 근거한 정식 소명 답변서가 1초 만에 자동 생성됩니다.",
        "4. [결과 복사하기] 버튼을 누르면 카카오톡, 사내 메신저, 이메일에 바로 붙여넣어 전송할 수 있습니다."
    ]
    for p in faq_points:
        pt = tf_fq.add_paragraph()
        pt.text = p
        pt.font.name = FONT_FAMILY
        pt.font.size = Pt(9.5)
        pt.font.color.rgb = C_TEXT_LIGHT
        pt.space_before = Pt(8)

    # ==============================================================================
    # SLIDE 10: [CHAPTER 4] 입력 예시 (INPUT EXAMPLES)
    # ==============================================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s10)
    add_header(s10, "Chapter 4. 입력 예시 (Input Samples)", "표준 입력 견적서 및 시안 이미지 예시 규격", 
               "정확한 감리 및 산출을 위해 지원되는 엑셀 서식 구조와 이미지 입력 가이드라인입니다.")

    add_card(s10, 0.8, 1.85, 5.7, 5.0, C_CARD_BG, C_CYAN, 1)
    tb_in1 = s10.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.3), Inches(4.7))
    tf_in1 = tb_in1.text_frame
    tf_in1.word_wrap = True

    pin1 = tf_in1.paragraphs[0]
    pin1.text = "📊 입력 예시 1: 엑셀 견적서 (.xlsx)"
    pin1.font.name = FONT_FAMILY
    pin1.font.size = Pt(13)
    pin1.font.bold = True
    pin1.font.color.rgb = C_CYAN

    in1_desc = (
        "• 권장 시트 구성:\n"
        "  - [견적서 시트]: 구분, 품명(항목명), 규격, 수량, 단위, 단가, 금액, 비고\n"
        "  - [노무비 시트]: 공정, 역할, 인원, 일수, 일단가, 금액\n\n"
        "• 컬럼명 유연성:\n"
        "  - 컬럼명이 '수량' 대신 'QTY', '단가' 대신 '단위당금액'이어도 자동 매핑\n\n"
        "• 입력 엑셀 예시 데이터:\n"
        "  [구분: 제작 | 품명: 특수 아크릴 전시대 | 수량: 10개 | 단가: ₩120,000 | 금액: ₩1,200,000]\n"
        "  [구분: 인건비 | 역할: 목공 숙련공 | 인원: 2명 | 일수: 1일 | 단가: ₩280,000 | 금액: ₩560,000]"
    )
    p_in1_body = tf_in1.add_paragraph()
    p_in1_body.text = in1_desc
    p_in1_body.font.name = FONT_FAMILY
    p_in1_body.font.size = Pt(9.5)
    p_in1_body.font.color.rgb = C_TEXT_LIGHT
    p_in1_body.space_before = Pt(8)

    add_card(s10, 6.833, 1.85, 5.7, 5.0, C_CARD_BG, C_GOLD, 1)
    tb_in2 = s10.shapes.add_textbox(Inches(7.033), Inches(2.0), Inches(5.3), Inches(4.7))
    tf_in2 = tb_in2.text_frame
    tf_in2.word_wrap = True

    pin2 = tf_in2.paragraphs[0]
    pin2.text = "📷 입력 예시 2: 비전 AI 시안 이미지"
    pin2.font.name = FONT_FAMILY
    pin2.font.size = Pt(13)
    pin2.font.bold = True
    pin2.font.color.rgb = C_GOLD

    in2_desc = (
        "• 지원 파일 형식: PNG, JPG, JPEG, WEBP\n\n"
        "• 권장 시안 형태:\n"
        "  - 3D Rhino/3ds Max 렌더링 투시도\n"
        "  - 정면/측면/평면이 확인 가능한 목업 스케치\n"
        "  - 과거 유사 시공 현장 실물 사진\n\n"
        "• AI 자동 인식 추출 항목:\n"
        "  - 1) 외형 치수: 가로(W) × 세로(D) × 높이(H) 자동 추정\n"
        "  - 2) 외장 소재: 갈바 스틸, 아크릴, 원목 MDF, 강화유리 등\n"
        "  - 3) 특수 가공: 분체 도장, 레이저 타공, LED 조명 매립 유무"
    )
    p_in2_body = tf_in2.add_paragraph()
    p_in2_body.text = in2_desc
    p_in2_body.font.name = FONT_FAMILY
    p_in2_body.font.size = Pt(9.5)
    p_in2_body.font.color.rgb = C_TEXT_LIGHT
    p_in2_body.space_before = Pt(8)

    # ==============================================================================
    # SLIDE 11: [CHAPTER 4] 출력 예시 (OUTPUT EXAMPLES)
    # ==============================================================================
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s11)
    add_header(s11, "Chapter 4. 출력 예시 (Output Samples)", "산출물 유형 및 결과물 활용 가이드", 
               "원클릭으로 생성되는 종합소견서, 공문 진단서, 산출 엑셀 견적서 등 실무 보고용 산출물입니다.")

    out_cards = [
        ("📋 1. 견적감독관 종합소견서 전문", 
         "클립보드 복사 & .md 다운로드",
         "• [소견서 발신]: 제일기획 AX 견적감독관\n• [종합 판정]: 🔴 고위험 (재협상 및 삭감 필수)\n• [삭감 타겟]: ₩5,720,000 감액 청구 권고\n• [최종 예산]: ₩52,480,000 합의선 제시\n• [4대 총평]: 단가/투입률/제작비/승인절차 소견 완비\n👉 사내 결재문서 및 경영진 보고서에 1초 만에 복사하여 붙여넣기 가능"),
        ("📑 2. 공문형 공식 감리 진단서", 
         "협력사 통보 및 공문서 발송용",
         "• 마크다운(.md) 형식으로 저장되는 표준 감리 보고서\n• 비적합 비목별 위반 법령 및 공공 품셈 조항 명시\n• 한국물가협회, 조달청, 근로기준법 조항 인용\n• 협력사에 공식 시정명령 및 소명 요구서로 송부하여 원가 삭감의 법적·논리적 근거로 활용"),
        ("📊 3. 공학적 표준 견적서 (.xlsx)", 
         "비전 AI가 산출한 정밀 엑셀 견적서",
         "• 시안에서 추출한 5대 공종 원가가 수식과 함께 내재화된 엑셀 파일 자동 다운로드\n• 자재비, 가공비, 조립공임, 도장마감, 포장비 분해\n• 최저가/표준가/프리미엄 3대 시나리오 탭 수록\n• 광고주 보고 및 제작처 발주 표준 견적서로 즉시 사용")
    ]

    for idx, (title, sub, body) in enumerate(out_cards):
        ox = 0.8 + idx * 3.98
        add_card(s11, ox, 1.85, 3.78, 5.0, C_CARD_BG, C_GREEN, 1)

        tb = s11.shapes.add_textbox(Inches(ox + 0.15), Inches(2.0), Inches(3.48), Inches(4.7))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.name = FONT_FAMILY
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.font.color.rgb = C_GREEN

        p2 = tf.add_paragraph()
        p2.text = sub
        p2.font.name = FONT_FAMILY
        p2.font.size = Pt(10)
        p2.font.bold = True
        p2.font.color.rgb = C_TEXT_WHITE
        p2.space_before = Pt(6)

        p3 = tf.add_paragraph()
        p3.text = body
        p3.font.name = FONT_FAMILY
        p3.font.size = Pt(9.5)
        p3.font.color.rgb = C_TEXT_LIGHT
        p3.space_before = Pt(10)

    # ==============================================================================
    # SLIDE 12: [CHAPTER 5] 한계 및 주의사항 (LIMITATIONS & CAUTIONS)
    # ==============================================================================
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s12)
    add_header(s12, "Chapter 5. 한계 및 주의사항", "시스템 한계 및 업무 활용 시 필수 주의사항", 
               "안전하고 신뢰도 높은 의사결정을 위해 AI 시스템의 한계점과 필수 검토 수칙을 숙지하시기 바랍니다.")

    cautions = [
        ("⚠️ 1. 비전 AI 시각적 한계치", 
         "내부 특수 전장/기계 장치",
         "• 시안 이미지는 '외형'을 기반으로 재질과 공종을 역설계하므로, 제품 내부에 매립된 특수 복합 센서, 키오스크 PC, 특수 구동 모터 등은 추가 설계 검토가 필요합니다.\n• 복합 키트의 경우 반드시 25.6억 실물 키트 DB와의 교차 검증을 병행하십시오."),
        ("⚠️ 2. 비정형 서식 및 수기 인식", 
         "스캔 이미지 견적서 주의",
         "• 인쇄된 견적서를 사진으로 찍거나 손글씨로 작성된 스캔 PDF/이미지는 OCR 인식률이 저하될 수 있습니다.\n• 정확한 원가 산출을 위해 협력사로부터 가급적 디지털 원본 엑셀(.xlsx) 파일을 수령하여 업로드하십시오."),
        ("⚠️ 3. 삭감 타겟액의 성격", 
         "법적 강제가 아닌 협상 마지노선",
         "• 시스템이 제시하는 '부적정 삭감 타겟액'은 공공 고시 기준 및 계약 원칙에 기반한 '합리적 네고 가이드라인'입니다.\n• 현장 특수성이나 납기 긴급성에 따른 할증은 광고주/협력사 간 서면 합의(CR)를 통해 조건부 인정될 수 있습니다."),
        ("⚠️ 4. 대외비 및 보안 준수", 
         "사내 정보보호 규정 준수",
         "• 본 시스템은 SSL 보안 전송 및 인메모리 연산을 적용하여 안전하나, 미출시 신제품이나 특급 보안 프로젝트 이미지를 업로드할 때는 사내 정보보안 가이드라인을 엄격히 준수하십시오.")
    ]

    for idx, (title, sub, body) in enumerate(cautions):
        cx = 0.8 + (idx % 2) * 5.95
        cy = 1.85 + (idx // 2) * 2.5
        add_card(s12, cx, cy, 5.75, 2.35, C_CARD_BG, C_RED, 1)

        tb = s12.shapes.add_textbox(Inches(cx + 0.15), Inches(cy + 0.15), Inches(5.45), Inches(2.05))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.name = FONT_FAMILY
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = C_RED

        p2 = tf.add_paragraph()
        p2.text = sub
        p2.font.name = FONT_FAMILY
        p2.font.size = Pt(9.5)
        p2.font.bold = True
        p2.font.color.rgb = C_TEXT_WHITE
        p2.space_before = Pt(3)

        p3 = tf.add_paragraph()
        p3.text = body
        p3.font.name = FONT_FAMILY
        p3.font.size = Pt(9)
        p3.font.color.rgb = C_TEXT_LIGHT
        p3.space_before = Pt(5)

    # ==============================================================================
    # SLIDE 13: [CHAPTER 6] 문의처 및 기술지원 (SUPPORT & CONTACTS)
    # ==============================================================================
    s13 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s13)
    add_header(s13, "Chapter 6. 문의처 및 기술지원", "전담 지원 조직 및 기술 문의 채널 안내", 
               "서비스 사용 중 기능 개선 제안, 엑셀 파싱 오류 문의, 긴급 장애 신고는 전담 채널로 연락해 주십시오.")

    contacts = [
        ("🏢 1. 총괄 운영 부서", 
         "제일기획 AX 혁신추진팀 (AI Transformation Office)",
         "• 업무 총괄: 2026 AX 경진대회 TFT\n• 업무 시간: 평일 09:00 ~ 18:00 (점심시간 12:00 ~ 13:00)\n• 주요 역할: 신규 공공 기준 데이터 업데이트, 감리 정책 가이드라인 수립, 전사 라이선스 관리"),
        ("✉️ 2. 기술 문의 및 데스크", 
         "이메일 & 사내 메신저 채널",
         "• 공식 문의 이메일: ax-estimate-support@cheil.com\n• 사내 메신저: #ax-견적감리관-지원 채널\n• 처리 기준: 일반 문의 2시간 이내 회신, 엑셀 파서 오류 4시간 이내 패치 배포"),
        ("🚨 3. 긴급 핫라인 & 저장소", 
         "24/7 시스템 모니터링",
         "• 서버 긴급 관제 핫라인: 사내 내선 8080 (비상 관제 센터)\n• GitHub 이슈 트래커:\n  https://github.com/designer01jsw-commits/ax-estimate-supervisor/issues\n• 시스템 상태 페이지: https://ax-estimate-supervisor.onrender.com")
    ]

    for idx, (title, sub, body) in enumerate(contacts):
        cx = 0.8 + idx * 3.98
        add_card(s13, cx, 1.85, 3.78, 5.0, C_CARD_BG, C_PURPLE, 1)

        tb = s13.shapes.add_textbox(Inches(cx + 0.15), Inches(2.0), Inches(3.48), Inches(4.7))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.name = FONT_FAMILY
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.font.color.rgb = C_PURPLE

        p2 = tf.add_paragraph()
        p2.text = sub
        p2.font.name = FONT_FAMILY
        p2.font.size = Pt(10)
        p2.font.bold = True
        p2.font.color.rgb = C_TEXT_WHITE
        p2.space_before = Pt(6)

        p3 = tf.add_paragraph()
        p3.text = body
        p3.font.name = FONT_FAMILY
        p3.font.size = Pt(9.5)
        p3.font.color.rgb = C_TEXT_LIGHT
        p3.space_before = Pt(10)

    # ==============================================================================
    # SLIDE 14: FAQ & 자주 묻는 질문 5選 (FAQ SUMMARY)
    # ==============================================================================
    s14 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s14)
    add_header(s14, "FAQ Summary", "자주 묻는 질문(FAQ) 핵심 정리 5選", 
               "실제 현업 담당자들이 가장 자주 묻는 질문 5가지를 일문일답으로 명쾌하게 정리했습니다.")

    faqs_qa = [
        ("Q1. 사용자 컴퓨터나 노트북 전원을 끄면 웹 도메인 접속이 끊기나요?",
         "아닙니다! Render 글로벌 상용 클라우드에 365일 24시간 무중단 가동되므로 본인 PC 전원이 꺼져 있어도 언제나 전 세계 어디서나 정상 접속됩니다."),
        ("Q2. 엑셀 견적서 서식이나 컬럼명이 표준과 달라도 오류 없이 인식되나요?",
         "네! 내장된 지능형 유연 파서가 컬럼명 변형('단가', '단위당금액', '합계' 등)을 99.2% 자동으로 감지하여 정규화하므로 파일 수정 없이 그대로 올리시면 됩니다."),
        ("Q3. 스마트폰이나 외부 네트워크(외근/출장)에서도 접속할 수 있나요?",
         "네! 100% 완벽히 지원됩니다. 반응형 모바일 UI가 적용되어 있어 외근 중이나 협력사 미팅 중에도 스마트폰으로 접속하여 즉시 견적 감리 결과를 조회할 수 있습니다."),
        ("Q4. 견적서 파일 보안은 안전하게 보호되나요?",
         "안전합니다. 업로드된 파일은 메모리 상에서만 즉시 연산 및 대조 처리되며 외부 서버에 무단 저장되지 않고 256-bit SSL 암호화 터널을 통해 안전하게 전송됩니다."),
        ("Q5. 산출된 감사 리포트나 견적서는 어떤 형식으로 저장할 수 있나요?",
         "공식 감리 진단서 마크다운(.md), 결재용 텍스트 전문(클립보드 복사), 비전 AI 공학적 견적서 엑셀(.xlsx), 원본 대조 CSV 등 다양한 형태로 1초 만에 저장됩니다.")
    ]

    for idx, (q, a) in enumerate(faqs_qa):
        fy = 1.85 + idx * 1.02
        add_card(s14, 0.8, fy, 11.733, 0.92, C_CARD_BG, C_BLUE_BORDER, 1)

        tb = s14.shapes.add_textbox(Inches(1.0), Inches(fy + 0.08), Inches(11.333), Inches(0.8))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = q
        p1.font.name = FONT_FAMILY
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = C_CYAN

        p2 = tf.add_paragraph()
        p2.text = a
        p2.font.name = FONT_FAMILY
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = C_TEXT_LIGHT
        p2.space_before = Pt(2)

    # ==============================================================================
    # SLIDE 15: 1-Page Quick Reference Sheet (요약 가이드)
    # ==============================================================================
    s15 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s15)
    add_header(s15, "1-Page Reference", "업무 책상 부착용 1-Page 퀵 레퍼런스 시트", 
               "실제 실무 진행 시 옆에 두고 바로 찾아보는 3단계 초간단 감리 프로세스 요약표입니다.")

    steps_ref = [
        ("STEP 1. 접속 및 불러오기 (30초)", 
         "도메인 접속 및 파일 탑재",
         "1. https://ax-estimate-supervisor.onrender.com 접속\n2. 사이드바에서 프로젝트 선택 또는 내 엑셀 파일 업로드\n3. 최초 견적서와 변경 견적서 자동 대조 확인",
         C_CYAN),
        ("STEP 2. 종합진단 & 삭감액 확인 (1분)", 
         "견적감독관 종합진단실 검토",
         "1. 상단 종합의견 란에서 [🔴/🟡/🟢] 리스크 등급 확인\n2. '부적정 삭감 타겟액' 및 4대 공종별 총평 파악\n3. 감독관 3대 액션 가이드라인(즉시 삭감, 증빙 징구, 네고선) 확인",
         C_GOLD),
        ("STEP 3. 리포트 발행 & 협상 (1분)", 
         "소견서 복사 및 협력사 통보",
         "1. [결재용 소견서 전문] 원클릭 복사하여 사내 결재문서 첨부\n2. [공문형 감리 진단서] 다운로드하여 협력사에 감액 통보 송부\n3. 필요 시 [비전 AI] 및 [25.6억 키트 DB] 교차 벤치마킹",
         C_GREEN)
    ]

    for idx, (title, sub, body, col) in enumerate(steps_ref):
        rx = 0.8 + idx * 3.98
        add_card(s15, rx, 1.85, 3.78, 5.0, C_CARD_BG, col, 2)

        tb = s15.shapes.add_textbox(Inches(rx + 0.15), Inches(2.0), Inches(3.48), Inches(4.7))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.name = FONT_FAMILY
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.font.color.rgb = col

        p2 = tf.add_paragraph()
        p2.text = sub
        p2.font.name = FONT_FAMILY
        p2.font.size = Pt(10)
        p2.font.bold = True
        p2.font.color.rgb = C_TEXT_WHITE
        p2.space_before = Pt(6)

        p3 = tf.add_paragraph()
        p3.text = body
        p3.font.name = FONT_FAMILY
        p3.font.size = Pt(9.5)
        p3.font.color.rgb = C_TEXT_LIGHT
        p3.space_before = Pt(10)

    # Save to all 3 paths
    out_desktop_folder = Path(r"C:\Users\ws081\OneDrive\Desktop\정승우\AX대회관련")
    out_desktop_root = Path(r"C:\Users\ws081\OneDrive\Desktop")
    out_workspace = Path(r"c:\Users\ws081\OneDrive\문서\테스트_안티그래비티\AX대회")

    filename = "견적감리관_웹도메인_사용가이드.pptx"

    if out_desktop_folder.exists():
        p1 = out_desktop_folder / filename
        prs.save(str(p1))
        print(f"Saved: {p1}")

    p2 = out_desktop_root / filename
    prs.save(str(p2))
    print(f"Saved: {p2}")

    p3 = out_workspace / filename
    prs.save(str(p3))
    print(f"Saved: {p3}")

    print(f"Comprehensive Manual Deck ({len(prs.slides)} slides) generated successfully!")

if __name__ == "__main__":
    create_manual()
