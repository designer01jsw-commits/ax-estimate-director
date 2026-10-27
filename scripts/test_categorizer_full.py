# -*- coding: utf-8 -*-
import sys
from pathlib import Path
sys.path.append(str(Path(__file__).resolve().parent.parent))

import openpyxl
from auditor.loader import read_estimate_sheet
from auditor.categorizer import (
    decompose_estimate_by_category,
    compare_decomposed_estimates,
    audit_all_new_items,
    generate_cross_project_audit_report
)

p_base = Path("data/samples/기준견적서4_전시_광고주컨펌.xlsx")
p_rev = Path("data/samples/전시_홍대점_변경내역.xlsx")
wb_base = openpyxl.load_workbook(p_base, data_only=True)
wb_rev = openpyxl.load_workbook(p_rev, data_only=True)

sheet_b = [s for s in wb_base.sheetnames if "수원" in s][0]
sheet_r = wb_rev.sheetnames[0]

items_base, _, _, _ = read_estimate_sheet(wb_base[sheet_b], p_base.name, sheet_b)
items_rev, _, _, _ = read_estimate_sheet(wb_rev[sheet_r], p_rev.name, sheet_r)

base_decomp = decompose_estimate_by_category(items_base)
rev_decomp = decompose_estimate_by_category(items_rev)
comp = compare_decomposed_estimates(base_decomp, rev_decomp)

audited_items, total_excess, fair_cnt, caution_cnt, excess_cnt = audit_all_new_items(
    items_rev, items_base, base_decomp, currency="KRW"
)

report = generate_cross_project_audit_report(
    "삼성 사업장 전시 기준 v4", p_base.name, p_rev.name,
    base_decomp, rev_decomp, comp,
    total_excess, fair_cnt, caution_cnt, excess_cnt, currency="KRW"
)

print("PASS! Total excess:", total_excess)
print(f"Fair: {fair_cnt}, Caution: {caution_cnt}, Excess: {excess_cnt}")
print("Report preview line length:", len(report.splitlines()))
