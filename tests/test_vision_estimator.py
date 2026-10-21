# -*- coding: utf-8 -*-
import io
import pytest
from PIL import Image
import openpyxl

from auditor.vision_estimator import (
    analyze_image_features,
    estimate_costs_from_vision,
    export_estimate_to_excel_bytes,
    convert_vision_estimate_to_auditor_items
)

def create_dummy_image_bytes(width=800, height=600, color=(100, 150, 200)):
    img = Image.new("RGB", (width, height), color=color)
    buf = io.BytesIO()
    img.save(buf, format="JPEG")
    return buf.getvalue()

def test_analyze_image_features():
    img_bytes = create_dummy_image_bytes(1200, 800)
    features = analyze_image_features(img_bytes, hints="삼성 갤럭시 스튜디오 홍대점 팝업스토어")
    assert features["success"] is True
    assert features["width"] == 1200
    assert features["height"] == 800
    assert "팝업스토어" in features["detected_type"] or "전시" in features["detected_type"]
    assert len(features["detected_materials"]) > 0

def test_estimate_costs_from_vision_default():
    img_bytes = create_dummy_image_bytes(800, 800)
    features = analyze_image_features(img_bytes)
    
    result = estimate_costs_from_vision(
        features=features,
        project_type="스마트 체험 키트 / 원형 개발",
        scale="중형 (5~15평 / 키트 3~5대)",
        tier="표준형 (Standard)",
        options={"야간할증_포함": False, "예비버퍼_추가": True, "SW개발_포함": True, "무선시건_포함": True}
    )
    
    assert "items" in result
    assert "summary" in result
    items = result["items"]
    assert len(items) >= 10
    
    # Check that categories exist
    cats = {it["구분"] for it in items}
    assert "I. 기획·디자인 원형개발" in cats
    assert "II. 기구제작·자재비" in cats
    assert "V. 제경비 및 기술료" in cats
    
    summary = result["summary"]
    assert summary["total_cost"] > 0
    assert summary["total_md"] > 0
    assert summary["scenario_comparison"]["budget"] < summary["total_cost"] < summary["scenario_comparison"]["premium"]

def test_estimate_costs_options_night_and_premium():
    img_bytes = create_dummy_image_bytes(800, 800)
    features = analyze_image_features(img_bytes)
    
    res_std = estimate_costs_from_vision(features, tier="표준형 (Standard)", options={"야간할증_포함": False})
    res_prem = estimate_costs_from_vision(features, tier="고급형 (Premium)", options={"야간할증_포함": True})
    
    assert res_prem["summary"]["total_cost"] > res_std["summary"]["total_cost"]

def test_export_estimate_to_excel_bytes():
    img_bytes = create_dummy_image_bytes(800, 800)
    features = analyze_image_features(img_bytes)
    est = estimate_costs_from_vision(features)
    
    xlsx_bytes = export_estimate_to_excel_bytes(est, "테스트 전시 견적서")
    assert len(xlsx_bytes) > 1000
    
    # Read back with openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    assert "예상견적서" in wb.sheetnames
    ws = wb["예상견적서"]
    assert "제일기획 AX 견적감독관" in str(ws["A1"].value)
    assert ws.max_row >= 12

def test_convert_vision_estimate_to_auditor_items():
    img_bytes = create_dummy_image_bytes(800, 800)
    features = analyze_image_features(img_bytes)
    est = estimate_costs_from_vision(features)
    
    auditor_items = convert_vision_estimate_to_auditor_items(est)
    assert len(auditor_items) == len(est["items"])
    for it in auditor_items:
        assert "항목명" in it
        assert "단가" in it
        assert "금액" in it
        assert it["금액"] == it["수량"] * it["단가"] or it["금액"] > 0

def test_real_kit_presets_and_images():
    from auditor.vision_estimator import SAMPLE_PRESETS, load_preset_image_bytes, find_ground_truth_kit
    
    assert "M02_CXP" in SAMPLE_PRESETS
    assert "M03_200MP" in SAMPLE_PRESETS
    assert "F03_FLIP" in SAMPLE_PRESETS
    
    img_bytes = load_preset_image_bytes("M02_CXP")
    assert len(img_bytes) > 1000
    assert img_bytes[:8] == b"\x89PNG\r\n\x1a\n"

    gt = find_ground_truth_kit(preset_key="M02_CXP")
    assert gt is not None
    assert gt["kit_id"] == "M02"
    assert gt["settled_amount"] == 214924059
    assert gt["prototype_cost"] == 117724059

def test_estimate_with_real_kit_ground_truth():
    img_bytes = create_dummy_image_bytes(800, 800)
    features = analyze_image_features(img_bytes, hints="43인치 디스플레이 CNC 머시닝 우레탄 2색 도장")
    est = estimate_costs_from_vision(features, preset_key="M02_CXP")
    
    assert "ground_truth" in est
    gt = est["ground_truth"]
    assert gt is not None
    assert gt["matched"] is True
    assert gt["kit_id"] == "M02"
    assert gt["settled_amount"] > 0

