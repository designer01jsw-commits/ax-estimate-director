# -*- coding: utf-8 -*-
"""auditor/loader.py 단위 테스트"""

from pathlib import Path
import pytest
import openpyxl

from auditor.errors import ValidationError
from auditor.loader import (
    parse_number,
    load_workbook_safe,
    resolve_sheet,
    read_estimate_sheet,
    clean_col_name
)

DATA_DIR = Path(__file__).resolve().parent.parent / "data"
SAMPLES_DIR = DATA_DIR / "samples"
BROKEN_DIR = SAMPLES_DIR / "broken"

def test_clean_col_name_aliases():
    """열 별칭 정규화 확인"""
    assert clean_col_name("  품명  ") == "항목명"
    assert clean_col_name("단가(원)") == "단가"
    assert clean_col_name("코드") == "항목코드"
    assert clean_col_name("분류") == "구분"
    assert clean_col_name("합계금액") == "금액"

def test_parse_number():
    """숫자 변환 확인: 쉼표, 원, 통화기호"""
    assert parse_number("1,200,000") == 1200000.0
    assert parse_number("₩85,000") == 85000.0
    assert parse_number("85,000원") == 85000.0
    assert parse_number("$1,200") == 1200.0
    assert parse_number("  45,000  ") == 45000.0
    assert parse_number("TBD") is None
    assert parse_number("") is None
    assert parse_number(None) is None

def test_p001_loader_features():
    """P001 파일 로드 검증: 3행 머리글 탐색, 합계 행 건너뛰기, 원본 행 일치, 산출식 불일치"""
    wb = load_workbook_safe(SAMPLES_DIR / "P001_최초견적.xlsx", "P001_최초견적.xlsx")
    ws, sname, _ = resolve_sheet(wb, "견적", "P001_최초견적.xlsx")
    items, mismatches, dup_codes, warnings = read_estimate_sheet(ws, "P001_최초견적.xlsx", sname)

    # 1. 항목 개수 확인 (8개 정상 항목, 합계 행 건너뜀)
    assert len(items) == 8
    
    # 2. 합계 행 건너뜀 확인
    assert any("건너뛰었습니다" in w for w in warnings)

    # 3. 산출식 불일치 확인 (E-003 그래픽 패널)
    assert len(mismatches) == 1
    assert mismatches[0]["code"] == "E-003"
    assert mismatches[0]["file_amount"] == 1850000.0
    assert mismatches[0]["calc_amount"] == 1800000.0

    # 4. 원본행 일치 확인: E-001은 엑셀 4행
    assert items[0]["항목코드"] == "E-001"
    assert items[0]["원본행"] == 4
    assert items[0]["row_key"] == "P001_최초견적.xlsx|견적|4"

def test_p002_tbd_warning():
    """P002 USD 파일에서 TBD 변환 실패 경고 및 제외 확인"""
    wb = load_workbook_safe(SAMPLES_DIR / "P002_최초견적.xlsx", "P002_최초견적.xlsx")
    ws, sname, _ = resolve_sheet(wb, "견적", "P002_최초견적.xlsx")
    items, mismatches, dup_codes, warnings = read_estimate_sheet(ws, "P002_최초견적.xlsx", sname)

    # U-04는 TBD로 인해 계산에서 제외되어 3개만 로드됨
    assert len(items) == 3
    assert any("숫자 변환 실패" in w for w in warnings)

def test_broken_fake_excel():
    """가짜 엑셀 파일 로드 시 한국어 안내 메시지 확인"""
    with pytest.raises(ValidationError) as exc_info:
        load_workbook_safe(BROKEN_DIR / "fake_excel.xlsx", "fake_excel.xlsx")
    assert "‘fake_excel.xlsx’는 엑셀(.xlsx) 파일이 아닙니다" in str(exc_info.value)

def test_broken_missing_sheet():
    """필수 시트 누락 시 한국어 안내 메시지 확인"""
    wb = load_workbook_safe(BROKEN_DIR / "wrong_sheet_name.xlsx", "wrong_sheet_name.xlsx")
    with pytest.raises(ValidationError) as exc_info:
        resolve_sheet(wb, "견적", "wrong_sheet_name.xlsx")
    assert "‘견적’ 시트를 찾지 못했습니다" in str(exc_info.value)
    assert "시트1, 요약" in str(exc_info.value)

def test_broken_missing_columns():
    """필수 열 누락 시 한국어 안내 메시지 확인"""
    wb = load_workbook_safe(BROKEN_DIR / "missing_columns.xlsx", "missing_columns.xlsx")
    ws, sname, _ = resolve_sheet(wb, "견적", "missing_columns.xlsx")
    with pytest.raises(ValidationError) as exc_info:
        read_estimate_sheet(ws, "missing_columns.xlsx", sname)
    assert "‘단가’ 열을 찾지 못했습니다" in str(exc_info.value)
